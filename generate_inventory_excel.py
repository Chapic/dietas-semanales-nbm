#!/usr/bin/env python3
"""Generate an inventory workbook with formulas, validations and a JSON export macro.

The script builds the spreadsheet with openpyxl and then post-processes the
archive to add an Excel 4 macro sheet plus a shape button wired to the macro.
"""
from __future__ import annotations

import tempfile
import xml.etree.ElementTree as ET
import zipfile
from datetime import date
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, NamedStyle, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


INVENTORY_HEADERS: Tuple[str, ...] = (
    "Nom article",
    "Ubicació",
    "Estoc actual",
    "Data prevista de consum",
    "Quantitat a consumir",
    "Estoc insuficient?",
    "Data límit per fer comanda",
    "Data d'arribada esperada",
    "Quantitat a demanar",
)

LOCATIONS = (
    "ARMARI BLANC MENJADOR",
    "CONGELADOR BUGADERIA",
    "CONGELADOR DÚPLEX",
    "NEVERA CUINA",
    "ARMARIS BLANCS MENJADOR",
)

SAMPLE_ROWS: List[Dict[str, object]] = [
    {
        "Nom article": "Carn picada",
        "Ubicació": "CONGELADOR DÚPLEX",
        "Estoc actual": 4.0,
        "Data prevista de consum": date(2024, 12, 5),
        "Quantitat a consumir": 3.5,
        "Data d'arribada esperada": date(2024, 12, 2),
        "Quantitat a demanar": 2.5,
    },
    {
        "Nom article": "Salmó fresc",
        "Ubicació": "NEVERA CUINA",
        "Estoc actual": 2.0,
        "Data prevista de consum": date(2024, 12, 8),
        "Quantitat a consumir": 2.0,
        "Data d'arribada esperada": date(2024, 12, 7),
        "Quantitat a demanar": 4.0,
    },
    {
        "Nom article": "Verdures variades",
        "Ubicació": "ARMARIS BLANCS MENJADOR",
        "Estoc actual": 6.0,
        "Data prevista de consum": date(2024, 12, 11),
        "Quantitat a consumir": 4.0,
        "Data d'arribada esperada": date(2024, 12, 9),
        "Quantitat a demanar": 1.0,
    },
    {
        "Nom article": "Pollastre sencer",
        "Ubicació": "CONGELADOR BUGADERIA",
        "Estoc actual": 1.0,
        "Data prevista de consum": date(2024, 12, 3),
        "Quantitat a consumir": 2.0,
        "Data d'arribada esperada": date(2024, 12, 4),
        "Quantitat a demanar": 3.0,
    },
]


def build_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventari"
    ws.freeze_panes = "A2"

    header_style = NamedStyle(name="header")
    header_style.font = Font(bold=True, color="FFFFFF")
    header_style.fill = PatternFill("solid", fgColor="44546A")
    header_style.alignment = Alignment(horizontal="center", vertical="center")

    number_style = NamedStyle(name="two_decimals")
    number_style.number_format = "0.00"
    number_style.alignment = Alignment(horizontal="right")

    date_style = NamedStyle(name="iso_date")
    date_style.number_format = "yyyy-mm-dd"
    date_style.alignment = Alignment(horizontal="center")

    order_deadline_style = NamedStyle(name="order_deadline")
    order_deadline_style.number_format = "yyyy-mm-dd"
    order_deadline_style.font = Font(color="F4A261")
    order_deadline_style.alignment = Alignment(horizontal="center")

    for style in (header_style, number_style, date_style, order_deadline_style):
        if style.name not in wb.named_styles:
            wb.add_named_style(style)

    ws.append(list(INVENTORY_HEADERS))

    for idx, header in enumerate(INVENTORY_HEADERS, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.style = "header"
        ws.column_dimensions[get_column_letter(idx)].width = 22

    for offset, row in enumerate(SAMPLE_ROWS, start=2):
        ws.cell(row=offset, column=1, value=row["Nom article"])
        ws.cell(row=offset, column=2, value=row["Ubicació"])
        ws.cell(row=offset, column=3, value=row["Estoc actual"]).style = "two_decimals"
        ws.cell(row=offset, column=4, value=row["Data prevista de consum"]).style = "iso_date"
        ws.cell(row=offset, column=5, value=row["Quantitat a consumir"]).style = "two_decimals"
        ws.cell(row=offset, column=6).value = f"=IF(C{offset}<E{offset},\"Sí\",\"No\")"
        ws.cell(row=offset, column=7).value = f"=D{offset}-15"
        ws.cell(row=offset, column=7).style = "order_deadline"
        ws.cell(row=offset, column=8, value=row["Data d'arribada esperada"]).style = "iso_date"
        ws.cell(row=offset, column=9, value=row["Quantitat a demanar"]).style = "two_decimals"

    # Extend styles downwards for blank template rows
    for row in range(2, 102):
        for col in (3, 5, 9):
            ws.cell(row=row, column=col).style = "two_decimals"
        for col in (4, 8):
            ws.cell(row=row, column=col).style = "iso_date"
        ws.cell(row=row, column=7).style = "order_deadline"

    location_validation = DataValidation(
        type="list",
        formula1='"' + ",".join(LOCATIONS) + '"',
        allow_blank=True,
        showDropDown=True,
    )
    ws.add_data_validation(location_validation)
    location_validation.add("B2:B500")

    decimal_validation = DataValidation(
        type="decimal",
        operator="greaterThanOrEqual",
        formula1="0",
        allow_blank=True,
    )
    ws.add_data_validation(decimal_validation)
    for col in ("C", "E", "I"):
        decimal_validation.add(f"{col}2:{col}500")

    date_validation = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)")
    ws.add_data_validation(date_validation)
    for col in ("D", "H"):
        date_validation.add(f"{col}2:{col}500")

    low_stock_rule = FormulaRule(formula=["$F2=\"Sí\""], fill=PatternFill("solid", fgColor="FFDCDC"))
    ok_stock_rule = FormulaRule(formula=["$F2=\"No\""], fill=PatternFill("solid", fgColor="DCFFDC"))
    ws.conditional_formatting.add("A2:I500", low_stock_rule)
    ws.conditional_formatting.add("A2:I500", ok_stock_rule)

    ws.cell(row=1, column=11, value="Utilitza el botó per exportar a JSON.")

    wb.save(path)


def _next_rid(existing: Iterable[str]) -> str:
    max_id = 0
    for rid in existing:
        if rid.startswith("rId"):
            try:
                max_id = max(max_id, int(rid[3:]))
            except ValueError:
                continue
    return f"rId{max_id + 1}"


def _write_macrosheet(macrosheet_path: Path) -> None:
    formulas = [
        'SET.NAME("firstDataRow",2)',
        'SET.NAME("lastRow",MATCH("zzz",Inventari!$A:$A))',
        'SET.NAME("rowCount",IF(ISNUMBER(lastRow),lastRow-firstDataRow+1,0))',
        'IF(rowCount<=0,RETURN(),)',
        'SET.NAME("jsonPath",CONCATENATE(GET.DOCUMENT(63),"inventory_export.json"))',
        'SET.NAME("fh",FOPEN(jsonPath,2))',
        'FWRITE(fh,"[")',
        'SET.NAME("r",0)',
        'LABEL(loopStart)',
        'SET.NAME("r",r+1)',
        'IF(r>rowCount,GOTO(closeSection),)',
        'SET.NAME("rowIndex",firstDataRow+r-1)',
        'SET.NAME("rowJson",CONCATENATE("{\"Nom article\":\"",INDEX(Inventari!$A:$A,rowIndex),"\",\"Ubicació\":\"",INDEX(Inventari!$B:$B,rowIndex),"\",\"Estoc actual\":",TEXT(INDEX(Inventari!$C:$C,rowIndex),"0.00"),",\"Data prevista de consum\":\"",TEXT(INDEX(Inventari!$D:$D,rowIndex),"yyyy-mm-dd"),"\",\"Quantitat a consumir\":",TEXT(INDEX(Inventari!$E:$E,rowIndex),"0.00"),",\"Estoc insuficient?\":\"",INDEX(Inventari!$F:$F,rowIndex),"\",\"Data límit per fer comanda\":\"",TEXT(INDEX(Inventari!$G:$G,rowIndex),"yyyy-mm-dd"),"\",\"Data d'arribada esperada\":\"",TEXT(INDEX(Inventari!$H:$H,rowIndex),"yyyy-mm-dd"),"\",\"Quantitat a demanar\":",TEXT(INDEX(Inventari!$I:$I,rowIndex),"0.00"),"}"))',
        'IF(r=rowCount,FWRITE(fh,rowJson),FWRITE(fh,CONCATENATE(rowJson,",")))',
        'GOTO(loopStart)',
        'LABEL(closeSection)',
        'FWRITE(fh,"]")',
        'FCLOSE(fh)',
        'ALERT(CONCATENATE("Exportació JSON creada a ",jsonPath))',
        'RETURN()'
    ]

    rows_xml = []
    for idx, formula in enumerate(formulas, start=1):
        cell = f"A{idx}"
        rows_xml.append(
            f'<row r="{idx}" spans="1:1"><c r="{cell}" t="str"><f>{formula}</f><v>0</v></c></row>'
        )

    macrosheet_xml = f"""<?xml version='1.0' encoding='UTF-8'?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <sheetPr codeName="ExportJSON"/>
  <dimension ref="A1:A{len(formulas)}"/>
  <sheetViews/>
  <sheetFormatPr defaultRowHeight="15"/>
  <cols>
    <col min="1" max="1" width="150" customWidth="1"/>
  </cols>
  <sheetData>
    {''.join(rows_xml)}
  </sheetData>
  <pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>
</worksheet>
"""
    macrosheet_path.write_text(macrosheet_xml, encoding="utf-8")


def _write_button(drawing_path: Path) -> None:
    drawing_xml = """<?xml version='1.0' encoding='UTF-8'?>
<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <xdr:twoCellAnchor editAs="oneCell">
    <xdr:from>
      <xdr:col>8</xdr:col>
      <xdr:colOff>0</xdr:colOff>
      <xdr:row>0</xdr:row>
      <xdr:rowOff>0</xdr:rowOff>
    </xdr:from>
    <xdr:to>
      <xdr:col>10</xdr:col>
      <xdr:colOff>0</xdr:colOff>
      <xdr:row>3</xdr:row>
      <xdr:rowOff>0</xdr:rowOff>
    </xdr:to>
    <xdr:sp macro="ExportInventoryJSON">
      <xdr:nvSpPr>
        <xdr:cNvPr id="2" name="ExportarJSON"/>
        <xdr:cNvSpPr/>
      </xdr:nvSpPr>
      <xdr:spPr>
        <a:solidFill>
          <a:srgbClr val="1D7874"/>
        </a:solidFill>
        <a:ln w="9525">
          <a:solidFill>
            <a:srgbClr val="FFFFFF"/>
          </a:solidFill>
        </a:ln>
      </xdr:spPr>
      <xdr:txBody>
        <a:bodyPr/>
        <a:lstStyle/>
        <a:p>
          <a:r>
            <a:rPr lang="ca-ES" sz="1200" b="1"/>
            <a:t>Exportar a JSON</a:t>
          </a:r>
          <a:endParaRPr lang="ca-ES"/>
        </a:p>
      </xdr:txBody>
    </xdr:sp>
    <xdr:clientData/>
  </xdr:twoCellAnchor>
</xdr:wsDr>
"""
    drawing_path.write_text(drawing_xml, encoding="utf-8")


def embed_macro(xlsx_path: Path, output_path: Path) -> None:
    ns_pkg = "http://schemas.openxmlformats.org/package/2006/content-types"
    ns_main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ns_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ns_pkg_rel = "http://schemas.openxmlformats.org/package/2006/relationships"

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        with zipfile.ZipFile(xlsx_path) as zf:
            zf.extractall(tmp_path)

        # Content types
        content_types = tmp_path / "[Content_Types].xml"
        tree = ET.parse(content_types)
        root = tree.getroot()
        workbook_override = root.find(f"{{{ns_pkg}}}Override[@PartName='/xl/workbook.xml']")
        if workbook_override is None:
            raise RuntimeError("Workbook override missing in content types")
        overrides = {(node.get("PartName"), node) for node in root.findall(f"{{{ns_pkg}}}Override")}
        needed_overrides = {
            "/xl/macrosheets/sheet1.xml": "application/vnd.ms-excel.macrosheet+xml",
            "/xl/drawings/drawing1.xml": "application/vnd.openxmlformats-officedocument.drawing+xml",
        }
        for part, content_type in needed_overrides.items():
            if not root.findall(f"{{{ns_pkg}}}Override[@PartName='{part}']"):
                elem = ET.SubElement(root, f"{{{ns_pkg}}}Override")
                elem.set("PartName", part)
                elem.set("ContentType", content_type)
        tree.write(content_types, encoding="utf-8", xml_declaration=True)

        # Workbook relationships
        wb_rels_path = tmp_path / "xl/_rels/workbook.xml.rels"
        rels_tree = ET.parse(wb_rels_path)
        rels_root = rels_tree.getroot()
        rel_ids = [rel.get("Id") for rel in rels_root.findall(f"{{{ns_pkg_rel}}}Relationship")]
        macro_rel_id = _next_rid(rel_ids)
        macro_rel = ET.SubElement(
            rels_root,
            f"{{{ns_pkg_rel}}}Relationship",
            {
                "Id": macro_rel_id,
                "Type": "http://schemas.microsoft.com/office/2006/relationships/xlMacrosheet",
                "Target": "macrosheets/sheet1.xml",
            },
        )
        rels_tree.write(wb_rels_path, encoding="utf-8", xml_declaration=True)

        # Workbook sheets and defined names
        wb_xml_path = tmp_path / "xl/workbook.xml"
        wb_tree = ET.parse(wb_xml_path)
        wb_root = wb_tree.getroot()
        sheets_node = wb_root.find(f".//{{{ns_main}}}sheets")
        if sheets_node is None:
            raise RuntimeError("Workbook sheets node missing")
        sheet_ids = [int(sheet.get("sheetId")) for sheet in sheets_node.findall(f"{{{ns_main}}}sheet")]
        next_sheet_id = max(sheet_ids) + 1
        macro_sheet = ET.SubElement(
            sheets_node,
            f"{{{ns_main}}}sheet",
            {
                "name": "Export JSON",
                "sheetId": str(next_sheet_id),
                "state": "veryHidden",
                f"{{{ns_rel}}}id": macro_rel_id,
            },
        )

        defined_names = wb_root.find(f"{{{ns_main}}}definedNames")
        if defined_names is None:
            defined_names = ET.SubElement(wb_root, f"{{{ns_main}}}definedNames")
        macro_defined_name = ET.SubElement(
            defined_names,
            f"{{{ns_main}}}definedName",
            {
                "name": "ExportInventoryJSON",
                "comment": "Exporta les dades actuals a JSON",
                "hidden": "1",
            },
        )
        macro_defined_name.text = "'Export JSON'!$A$1"
        wb_tree.write(wb_xml_path, encoding="utf-8", xml_declaration=True)

        # Macro sheet file
        macrosheet_dir = tmp_path / "xl/macrosheets"
        macrosheet_dir.mkdir(exist_ok=True)
        _write_macrosheet(macrosheet_dir / "sheet1.xml")

        # Drawing relationships for the main sheet
        sheet_rels_dir = tmp_path / "xl/worksheets/_rels"
        sheet_rels_dir.mkdir(exist_ok=True)
        main_sheet_rels = sheet_rels_dir / "sheet1.xml.rels"
        if main_sheet_rels.exists():
            sheet_rels_tree = ET.parse(main_sheet_rels)
            sheet_rels_root = sheet_rels_tree.getroot()
        else:
            sheet_rels_root = ET.Element(f"{{{ns_pkg_rel}}}Relationships")
            sheet_rels_tree = ET.ElementTree(sheet_rels_root)
        sheet_rel_ids = [rel.get("Id") for rel in sheet_rels_root.findall(f"{{{ns_pkg_rel}}}Relationship")]
        drawing_rel_id = _next_rid(sheet_rel_ids)
        ET.SubElement(
            sheet_rels_root,
            f"{{{ns_pkg_rel}}}Relationship",
            {
                "Id": drawing_rel_id,
                "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing",
                "Target": "../drawings/drawing1.xml",
            },
        )
        sheet_rels_tree.write(main_sheet_rels, encoding="utf-8", xml_declaration=True)

        # Add drawing reference to the worksheet
        sheet_xml_path = tmp_path / "xl/worksheets/sheet1.xml"
        sheet_tree = ET.parse(sheet_xml_path)
        sheet_root = sheet_tree.getroot()
        drawing_node = sheet_root.find(f"{{{ns_main}}}drawing")
        if drawing_node is None:
            drawing_node = ET.Element(f"{{{ns_main}}}drawing")
            sheet_root.append(drawing_node)
        drawing_node.set(f"{{{ns_rel}}}id", drawing_rel_id)
        sheet_tree.write(sheet_xml_path, encoding="utf-8", xml_declaration=True)

        # Drawing file
        drawings_dir = tmp_path / "xl/drawings"
        drawings_dir.mkdir(exist_ok=True)
        _write_button(drawings_dir / "drawing1.xml")

        # Repack archive
        with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as zf:
            for fs_path in sorted(tmp_path.rglob("*")):
                if fs_path.is_file():
                    zf.write(fs_path, fs_path.relative_to(tmp_path))


def main() -> None:
    output = Path("inventari.xlsx")
    temp_xlsx = Path("_temp_inventory.xlsx")
    build_workbook(temp_xlsx)
    embed_macro(temp_xlsx, output)
    temp_xlsx.unlink(missing_ok=True)
    print(f"Fitxer generat: {output.resolve()}")


if __name__ == "__main__":
    main()
